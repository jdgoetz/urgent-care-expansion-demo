import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from reporting.generate_demo_report import DATA, build_report


class DemoReportTests(unittest.TestCase):
    def test_payload_is_public_demo_v2(self):
        payload = json.loads(DATA.read_text(encoding="utf-8"))
        self.assertEqual(payload["model"], "demo_expansion_score_v2")
        self.assertAlmostEqual(sum(payload["weights"].values()), 1.0)
        self.assertTrue(all(len(market["metrics"]) == 8 for market in payload["markets"]))

    def test_report_uses_six_column_breakdowns(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "demo.docx"
            build_report(output)
            document = Document(output)
            headers = [[cell.text for cell in table.rows[0].cells] for table in document.tables]
            self.assertIn(["Metric", "Raw Value", "Normalized Score", "Weight", "Contribution", "Explanation"], headers)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Demo Expansion Score v2", text)
            self.assertNotIn("production ranking", text.lower())


if __name__ == "__main__":
    unittest.main()
