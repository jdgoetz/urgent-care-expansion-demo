from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "reporting" / "public_report_data.json"
GEOMETRY = ROOT / "data" / "public" / "curated-zcta.json"
OUTPUT = ROOT / "docs" / "sample" / "urgent_care_expansion_demo_report.docx"
ASSET_DIR = ROOT / "reporting" / "output"
NAVY = "173047"
TEAL = "136F63"
BLUE = "1F6F8B"
LIGHT = "E8EEF2"
GRAY = "607183"


def shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    properties.append(node)


def add_table(document: Document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_properties.append(header_marker)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, LIGHT)
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True
        run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    return table


def add_map(market_name: str, zip_code: str, listed: bool) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / ("map_" + zip_code + ".png")
    image = Image.new("RGB", (900, 360), "#edf2f4")
    draw = ImageDraw.Draw(image)
    collection = json.loads(GEOMETRY.read_text(encoding="utf-8"))
    feature = next(item for item in collection["features"] if item["properties"]["ZCTA5"] == zip_code)
    geometry = feature["geometry"]
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    rings = [polygon[0] for polygon in polygons]
    points = [point for ring in rings for point in ring]
    min_x, max_x = min(point[0] for point in points), max(point[0] for point in points)
    min_y, max_y = min(point[1] for point in points), max(point[1] for point in points)
    scale = min(760 / max(max_x - min_x, 0.001), 250 / max(max_y - min_y, 0.001))
    fill = "#d9cbea" if listed else "#cfe5df"
    outline = "#6f4aa8" if listed else "#136f63"
    for ring in rings:
        rendered = [
            (450 + (point[0] - (min_x + max_x) / 2) * scale, 175 - (point[1] - (min_y + max_y) / 2) * scale)
            for point in ring
        ]
        draw.polygon(rendered, fill=fill, outline=outline, width=4)
    draw.text((30, 310), f"{market_name} | Census ZCTA {zip_code}", fill="#173047")
    image.save(path)
    return path


def configure(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)


def build_report():
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    document = Document()
    configure(document)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(100)
    title_run = title.add_run("Urgent Care Expansion\nIntelligence Demo")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph("Curated real-market public portfolio")
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(30)
    note.add_run("Boundary: ").bold = True
    note.add_run("Market geography and labeled public observations are real. Demo scores and sanitized scenarios are illustrative; production methodology is excluded.")
    document.add_page_break()

    document.add_heading("Executive Methodology", level=1)
    document.add_paragraph("A compact public demonstration of the decision workflow:")
    add_table(document, ["Stage", "Decision question"], [
        ["Market fundamentals", "Do public observations indicate structural attractiveness?"],
        ["Demo Expansion Score", "How do five transparent normalized components compare?"],
        ["Entry opportunities", "Is there a public or clearly illustrative near-term entry path?"],
        ["Entry Feasibility", "How strong is the best available demo path?"],
        ["Near-Term Priority", "Where should demo diligence happen first?"],
    ], [2.0, 5.0])
    document.add_paragraph()
    document.add_heading("Ranked Markets", level=2)
    add_table(document, ["Rank", "Market", "Metro", "Demo Near-Term", "Demo Expansion", "Demo Entry"], [
        [market["rank"], market["name"], market["metro"], market["near_term"], market["expansion"], market["entry"]]
        for market in payload["markets"]
    ])
    document.add_paragraph("Public model: demo_expansion_score_v1. This report does not reproduce production weights or opportunity logic.")

    for market in payload["markets"]:
        document.add_page_break()
        document.add_heading(str(market["rank"]) + ". " + market["name"], level=1)
        document.add_paragraph(market["metro"] + " | " + market["state"] + " | ZIP " + market["zip"])
        score_table = add_table(document, ["Demo Near-Term Priority", "Demo Expansion Score", "Demo Entry Feasibility"], [[market["near_term"], market["expansion"], market["entry"]]])
        for cell in score_table.rows[1].cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(15)
                run.font.color.rgb = RGBColor.from_string(TEAL)
        map_path = add_map(market["name"], market["zip"], market.get("listed", False))
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        map_image = paragraph.add_run().add_picture(str(map_path), width=Inches(5.9))
        map_image._inline.docPr.set("title", "Census ZCTA market geometry")
        map_image._inline.docPr.set("descr", "Census ZCTA geometry for " + market["name"])
        caption = document.add_paragraph("U.S. Census Bureau TIGERweb 2020 ZCTA geometry", style=None)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(8)

        document.add_heading("Market Overview", level=2)
        add_table(document, ["Indicator", "Value"], market["kpis"], [4.6, 2.4])
        document.add_heading("Competitors", level=2)
        add_table(document, ["Competitor", "Rating", "Reviews", "Weekly hours", "Operator / type"], market["competitors"])
        document.add_heading("Entry Opportunities", level=2)
        add_table(document, ["Path", "Opportunity / scenario", "Demo score", "Confidence"], market["opportunities"])
        document.add_heading("Metric Scoring", level=2)
        add_table(document, ["Metric", "Normalized", "Weight", "Contribution"], market["metrics"])
        document.add_paragraph("Source lineage: ACS 2019/2024 5-year estimates, Census TIGERweb geometry, linked public operator pages, and explicitly labeled public or illustrative opportunity evidence. No production records are included.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
